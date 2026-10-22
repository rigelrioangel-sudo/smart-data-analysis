"""Full workflow through the real HTTP API: upload → profile → review cleaning → analyze → export."""
import io
import json
import os
import sys

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
sys.path.insert(0, os.path.dirname(__file__))

import main  # noqa: E402
from sample_data import sample_csv_bytes  # noqa: E402

client = TestClient(main.app)


def upload(name, content):
    r = client.post("/upload", files={"file": (name, content)})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture()
def ds():
    return upload("sales.csv", sample_csv_bytes())


def approve_all(d):
    return {a["id"]: {"enabled": True} for a in d["plan"]}


def test_upload_returns_profile_plan_and_analysis(ds):
    assert len(ds["dataset_id"]) == 32
    assert ds["profile_original"]["row_count"] == 247
    assert ds["plan"] and ds["insights"] and ds["eda"]["relationships"]
    # Default view = safe fixes only: no rows removed, review items pending
    rep = ds["cleaning"]["report"]
    assert rep["rows"] == {"before": 247, "after": 247, "removed": 0}
    assert rep["pending_review"]
    assert rep["values"]["changed"] + rep["values"]["retained"] + rep["values"]["excluded"] == rep["values"]["total_before"]


def test_full_workflow_clean_analyze_export_undo_reset(ds):
    did = ds["dataset_id"]
    r = client.post(f"/datasets/{did}/clean", json={"actions": approve_all(ds)})
    assert r.status_code == 200, r.text
    d = r.json()
    rep = d["cleaning"]["report"]
    assert rep["rows"] == {"before": 247, "after": 240, "removed": 7}
    assert rep["duplicates"]["removed_exact"] == 5 and rep["duplicates"]["removed_by_key"] == 2
    assert rep["duplicates"]["remaining_exact"] == 0
    assert rep["invalid"]["detected"] == 4 and rep["invalid"]["unresolved"] == 0
    assert rep["columns"]["removed"] == ["currency", "notes"]
    v = rep["values"]
    assert v["changed"] + v["retained"] + v["excluded"] == v["total_before"]
    assert v["excluded"] == 7 * 10 + 240 * 2
    assert not rep["pending_review"]
    assert d["total_rows"] == 240 and d["can_undo"]

    # Cleaned CSV export matches the reported cleaned dataset
    csv_bytes = client.get(f"/datasets/{did}/export/cleaned", params={"format": "csv"}).content
    exported = pd.read_csv(io.BytesIO(csv_bytes), encoding="utf-8-sig")
    assert len(exported) == 240 and list(exported.columns) == d["columns"]
    assert exported["units"].isna().sum() == d["profile"]["columns"][d["columns"].index("units")]["missing_total"]

    # Excel and JSON exports load and agree
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/datasets/{did}/export/cleaned", params={"format": "xlsx"}).content))
    assert wb.active.max_row == 241
    records = client.get(f"/datasets/{did}/export/cleaned", params={"format": "json"}).json()
    assert len(records) == 240 and set(records[0]) == set(d["columns"])

    # Cleaning log contains every operation and the original values
    log = client.get(f"/datasets/{did}/export/cleaning-log", params={"format": "json"}).json()
    assert {o["action_id"] for o in log["operations"]} == {a["id"] for a in d["plan"]}
    assert any(c["before"] == "abc" and c["after"] is None for c in log["cell_changes"])
    assert len(log["excluded_rows"]) == 7
    ops_csv = pd.read_csv(io.BytesIO(client.get(f"/datasets/{did}/export/cleaning-log", params={"format": "csv"}).content))
    assert len(ops_csv) == len(d["plan"])
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/datasets/{did}/export/cleaning-log", params={"format": "xlsx"}).content))
    assert wb.sheetnames == ["Operations", "Cell changes", "Excluded rows", "Recipe"]

    # Reports
    qr = client.get(f"/datasets/{did}/export/quality-report", params={"format": "json"}).json()
    assert qr["summary"]["rows"]["after"] == 240 and len(qr["per_column"]) == 10
    html_q = client.get(f"/datasets/{did}/export/quality-report", params={"format": "html"}).text
    assert "Before and after cleaning" in html_q and "247 → 240" in html_q
    html_a = client.get(f"/datasets/{did}/export/analysis-report", params={"format": "html"})
    assert html_a.headers["content-type"].startswith("text/html")
    assert "<svg" in html_a.text and "Key insights" in html_a.text and "Limitations" in html_a.text
    ar = client.get(f"/datasets/{did}/export/analysis-report", params={"format": "json"}).json()
    assert ar["insights"] == d["insights"]
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/datasets/{did}/export/analysis-report", params={"format": "xlsx"}).content))
    assert "Insights" in wb.sheetnames and "Relationships" in wb.sheetnames
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/datasets/{did}/export/quality-report", params={"format": "xlsx"}).content))
    assert wb["Summary"]["B2"].value == 247 and wb["Summary"]["C2"].value == 240

    # Undo returns to the defaults, reset returns to the original
    r = client.post(f"/datasets/{did}/undo")
    assert r.status_code == 200 and r.json()["cleaning"]["report"]["rows"]["after"] == 247
    r = client.post(f"/datasets/{did}/reset")
    rep = r.json()["cleaning"]["report"]
    assert rep["values"]["changed"] == 0 and rep["rows"]["after"] == 247
    original = pd.read_csv(io.BytesIO(client.get(f"/datasets/{did}/export/cleaned", params={"format": "csv"}).content),
                           dtype=str, keep_default_na=False, encoding="utf-8-sig")
    source = pd.read_csv(io.BytesIO(sample_csv_bytes()), dtype=str, keep_default_na=False)
    pd.testing.assert_frame_equal(original, source)  # reset == the uploaded file, value for value


def test_recipe_reproduces_cleaned_output(ds):
    did = ds["dataset_id"]
    decisions = approve_all(ds)
    decisions["outliers:units"] = {"enabled": True, "options": {"handling": "set_missing", "method": "iqr1.5"}}
    client.post(f"/datasets/{did}/clean", json={"actions": decisions})
    recipe = client.get(f"/datasets/{did}/export/recipe", params={"format": "json"}).json()
    first = client.get(f"/datasets/{did}/export/cleaned", params={"format": "csv"}).content

    other = upload("sales.csv", sample_csv_bytes())
    r = client.post(f"/datasets/{other['dataset_id']}/clean", json=recipe)
    assert r.status_code == 200, r.text
    second = client.get(f"/datasets/{other['dataset_id']}/export/cleaned", params={"format": "csv"}).content
    assert first == second


def test_rows_search_sort_filter_and_versions(ds):
    did = ds["dataset_id"]
    rows = client.get(f"/datasets/{did}/rows", params={"version": "original", "search": "abc"}).json()
    assert rows["matched_rows"] == 2 and all("abc" in r["values"] for r in rows["rows"])

    filt = json.dumps([{"column": "revenue", "op": "gt", "value": "1000"}])
    r = client.get(f"/datasets/{did}/rows", params={"filters": filt}).json()
    assert r["matched_rows"] == 2 and {row["row"] for row in r["rows"]} == {81, 82}

    r = client.get(f"/datasets/{did}/rows", params={"sort": "revenue", "order": "desc", "page_size": 3}).json()
    vals = [float(row["values"][r["columns"].index("revenue")]) for row in r["rows"]]
    assert vals == sorted(vals, reverse=True) and vals[0] == 2400.5

    # Cleaned view marks changed cells with their original value
    r = client.get(f"/datasets/{did}/rows", params={"changed_only": "true", "search": "west", "page_size": 200}).json()
    region_idx = r["columns"].index("Region")
    trimmed = [row for row in r["rows"] if str(region_idx) in row["original"]]
    assert {row["row"] for row in trimmed} == {4, 8, 12}
    assert all(row["original"][str(region_idx)] == " West " and row["values"][region_idx] == "West" for row in trimmed)

    r = client.get(f"/datasets/{did}/rows", params={"filters": json.dumps([{"column": "units", "op": "is_missing"}])}).json()
    assert r["matched_rows"] == 4  # 'N/A' markers converted by the safe fix

    p2 = client.get(f"/datasets/{did}/rows", params={"page": 2, "page_size": 100}).json()
    assert p2["page"] == 2 and p2["rows"][0]["row"] == 101 and p2["total_pages"] == 3


def test_changes_endpoint(ds):
    did = ds["dataset_id"]
    r = client.get(f"/datasets/{did}/changes", params={"column": "Region"}).json()
    assert r["total"] == 3 and all(c["action"] == "trim_whitespace:Region" for c in r["changes"])


def test_datasets_are_isolated():
    a = upload("a.csv", b"x,y\n1,2\n3,4\n")
    b = upload("b.csv", b"name\nAlice\nBob\nCara\n")
    ra = client.get(f"/datasets/{a['dataset_id']}/rows").json()
    rb = client.get(f"/datasets/{b['dataset_id']}/rows").json()
    assert ra["columns"] == ["x", "y"] and rb["columns"] == ["name"] and rb["total_rows"] == 3
    assert client.delete(f"/datasets/{a['dataset_id']}").status_code == 200
    assert client.get(f"/datasets/{a['dataset_id']}").status_code == 404
    assert client.get(f"/datasets/{b['dataset_id']}").status_code == 200


@pytest.mark.parametrize("path,params,status", [
    ("/datasets/not-a-valid-id", {}, 422),
    ("/datasets/{id}/rows", {"filters": "not json"}, 400),
    ("/datasets/{id}/rows", {"filters": json.dumps([{"column": "nope", "op": "equals", "value": "1"}])}, 400),
    ("/datasets/{id}/rows", {"filters": json.dumps([{"column": "units", "op": "gt", "value": "abc"}])}, 400),
    ("/datasets/{id}/rows", {"sort": "nope"}, 400),
    ("/datasets/{id}/rows", {"page_size": 100000}, 422),
    ("/datasets/{id}/export/cleaned", {"format": "html"}, 400),
    ("/datasets/{id}/export/secrets", {"format": "csv"}, 422),
])
def test_input_validation(ds, path, params, status):
    r = client.get(path.replace("{id}", ds["dataset_id"]), params=params)
    assert r.status_code == status, r.text


def test_clean_rejects_bad_decisions(ds):
    did = ds["dataset_id"]
    r = client.post(f"/datasets/{did}/clean", json={"actions": {"outliers:units": {"enabled": True, "options": {"handling": "rm -rf"}}}})
    assert r.status_code == 400
    assert client.post(f"/datasets/{did}/clean", json={"format": "something-else", "actions": {}}).status_code == 400
    assert client.post(f"/datasets/{did}/clean", json={"actions": {"x": {"enabled": "maybe"}}}).status_code == 422


def test_unsafe_filename_and_formula_cells_are_neutralized():
    content = b"name,formula\nok,=HYPERLINK(\"http://x\")\nneg,-3\nat,@cmd\n"
    d = upload("../../etc/pa\"ss<wd>.csv", content)
    assert "/" not in d["filename"] and '"' not in d["filename"] and "<" not in d["filename"]
    did = d["dataset_id"]
    r = client.get(f"/datasets/{did}/export/cleaned", params={"format": "csv"})
    disp = r.headers["content-disposition"]
    assert disp.startswith('attachment; filename="') and ".." not in disp
    text = r.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'@cmd" in text and ",-3" in text
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/datasets/{did}/export/cleaned", params={"format": "xlsx"}).content))
    cell = wb.active["B2"]
    assert cell.value == '=HYPERLINK("http://x")' and cell.data_type == "s"  # text, not a formula


def test_upload_size_limit(monkeypatch):
    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 100)
    r = client.post("/upload", files={"file": ("big.csv", b"a\n" + b"1\n" * 200)})
    assert r.status_code == 413


@pytest.mark.parametrize("name,content", [
    ("mixed.csv", b"v\n1\ntwo\n3\n"),
    ("dates.csv", b"d\n2024-02-30\nnot a date\n2024-01-01\n"),
    ("allmissing.csv", b"a,b\n,\nN/A,\n"),
    ("single.csv", b"only\n1\n"),
    ("wide.json", json.dumps([{f"c{i}": i for i in range(60)}]).encode()),
    ("nested.json", json.dumps({"data": [{"id": 1, "tags": ["a"], "u": {"n": "x"}}]}).encode()),
])
def test_edge_case_files_complete_the_workflow(name, content):
    d = upload(name, content)
    did = d["dataset_id"]
    r = client.post(f"/datasets/{did}/clean", json={"actions": approve_all(d)})
    assert r.status_code == 200, r.text
    for art, fmt in [("cleaned", "xlsx"), ("analysis-report", "html"), ("quality-report", "xlsx"), ("cleaning-log", "csv")]:
        assert client.get(f"/datasets/{did}/export/{art}", params={"format": fmt}).status_code == 200
